from src.components import suggest_support_vehicle, parse_vehicle_capacity
from src.utils import format_number_it, format_date_it
from pathlib import Path
from typing import Optional, Union, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from datetime import datetime

from src.config import (
    PLANNER_TEMPLATE_FILE,
    OUTPUT_DIR
)


def _format_sga_date(val: Any) -> str:
    """Converte qualsiasi data nel formato rigido gg/mm/aaaa (DD/MM/YYYY)."""
    if pd.isna(val) or val is None or str(val).strip() == "":
        return ""
    s = str(val).strip().split(" ")[0]
    if "-" in s:
        parts = s.split("-")
        if len(parts) == 3 and len(parts[0]) == 4:
            return f"{parts[2]}/{parts[1]}/{parts[0]}"
    if "/" in s:
        parts = s.split("/")
        if len(parts) == 3 and len(parts[2]) == 4:
            return f"{int(parts[0]):02d}/{int(parts[1]):02d}/{parts[2]}"
    return s


def export_sga_file(df_flows: pd.DataFrame, output_path: Union[str, Path]):
    """
    Exports SGA flows (Barra 0 or Barra 9) to an Excel file and CSV matching rigid TMS requirements:
    Columns: GIRO, DATA_CARICO, COD_MITTENTE, MAGAZZINO_CARICO, DATA_SCARICO, COD_DESTINATARIO, MAGAZZINO_SCARICO, PESO, VOLUME, COLLI
    - Date strictly formatted as gg/mm/aaaa (DD/MM/YYYY).
    - Stores 'DA ASSEGNARE' / unassigned are strictly excluded.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    cols_order = [
        "GIRO", "DATA_CARICO", "COD_MITTENTE", "MAGAZZINO_CARICO",
        "DATA_SCARICO", "COD_DESTINATARIO", "MAGAZZINO_SCARICO",
        "PESO", "VOLUME", "COLLI"
    ]

    # Esclusione rigida negozi 'DA ASSEGNARE' o senza giro definito
    if df_flows is not None and not df_flows.empty:
        df_flows = df_flows.copy()
        # Normalizzazione nomi colonne (gestisce sia maiuscolo da engine che minuscolo da SQLite)
        col_rename = {}
        for c in df_flows.columns:
            u = str(c).strip().upper()
            if u in cols_order:
                col_rename[c] = u
            elif u in ("FLUSSO_TIPO", "STORE_CODE", "BARRA"):
                col_rename[c] = u.lower()
        if col_rename:
            df_flows = df_flows.rename(columns=col_rename)

        giro_col = "GIRO" if "GIRO" in df_flows.columns else ("giro" if "giro" in df_flows.columns else None)
        if giro_col:
            mask_valid = ~df_flows[giro_col].fillna("").astype(str).str.strip().str.upper().apply(
                lambda g: g in ("", "DA ASSEGNARE", "NON ASSEGNATO", "NON DEFINITO", "NONE", "NAN", "-")
                or g.startswith("DA ASS")
                or g.startswith("NON ASS")
            )
        else:
            mask_valid = pd.Series([True] * len(df_flows), index=df_flows.index)

        store_col = "store_code" if "store_code" in df_flows.columns else ("STORE_CODE" if "STORE_CODE" in df_flows.columns else None)
        if store_col:
            mask_valid = mask_valid & ~df_flows[store_col].fillna("").astype(str).str.strip().str.upper().apply(
                lambda s: s in ("DA ASSEGNARE", "NON ASSEGNATO") or s.startswith("DA ASS")
            )
        df_flows = df_flows[mask_valid].copy()
    else:
        df_flows = pd.DataFrame(columns=cols_order)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SGA_DEF"

    # Header style
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx, col_name in enumerate(cols_order, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for r_idx, (_, row) in enumerate(df_flows.iterrows(), start=2):
        flusso_val = str(row.get("flusso_tipo") or row.get("FLUSSO_TIPO") or "").strip().upper()
        is_ritorno = (flusso_val == "RITORNO")
        for c_idx, col_name in enumerate(cols_order, start=1):
            val = row.get(col_name)
            if val is None or pd.isna(val):
                val = row.get(col_name.lower())
            if is_ritorno and col_name in ("PESO", "VOLUME", "COLLI"):
                val = 0
            elif col_name in ("DATA_CARICO", "DATA_SCARICO"):
                val = _format_sga_date(val)
            elif val is None or pd.isna(val):
                val = 0 if col_name in ("PESO", "VOLUME", "COLLI") else ""
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if col_name in ("PESO", "VOLUME", "COLLI"):
                if is_ritorno:
                    cell.number_format = "0"
                else:
                    cell.number_format = "0.00" if col_name == "VOLUME" else "0"

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)

    # Save matching rigid CSV with DD/MM/YYYY dates
    csv_path = output_path.with_suffix(".csv")
    csv_rows = []
    for _, row in df_flows.iterrows():
        flusso_val = str(row.get("flusso_tipo") or row.get("FLUSSO_TIPO") or "").strip().upper()
        is_ritorno = (flusso_val == "RITORNO")
        csv_rows.append({
            "GIRO": row.get("GIRO", row.get("giro", "")),
            "DATA_CARICO": _format_sga_date(row.get("DATA_CARICO", row.get("data_carico", ""))),
            "COD_MITTENTE": row.get("COD_MITTENTE", row.get("cod_mittente", "")),
            "MAGAZZINO_CARICO": row.get("MAGAZZINO_CARICO", row.get("magazzino_carico", "")),
            "DATA_SCARICO": _format_sga_date(row.get("DATA_SCARICO", row.get("data_scarico", ""))),
            "COD_DESTINATARIO": row.get("COD_DESTINATARIO", row.get("cod_destinatario", "")),
            "MAGAZZINO_SCARICO": row.get("MAGAZZINO_SCARICO", row.get("magazzino_scarico", "")),
            "PESO": 0 if is_ritorno else (row.get("PESO") if row.get("PESO") is not None else row.get("peso", 0)),
            "VOLUME": 0 if is_ritorno else (row.get("VOLUME") if row.get("VOLUME") is not None else row.get("volume", 0)),
            "COLLI": 0 if is_ritorno else (row.get("COLLI") if row.get("COLLI") is not None else row.get("colli", 0)),
        })
    pd.DataFrame(csv_rows, columns=cols_order).to_csv(csv_path, index=False)


def export_zara_spunta(
    df_spunta: pd.DataFrame,
    output_path: Union[str, Path],
    special_breakdowns: Optional[Dict[str, Any]] = None
):
    """
    Exports the complete SPUNTA ARRIVI sheet as ZARA.xlsx,
    including the bottom section for Platform Processing (Lavorazioni di Piattaforma).
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SPUNTA ARRIVI "

    headers = [
        "NEGOZIO", "STORE", "PUNTO VENDITA", "MC", "TEXTILE ATTESI", "",
        "OTHER ATTESI", "", "SHOES ATTESI", "", "ECOM ATTESI", "",
        "PEDANE", "TOTALE PEDANE", "VETTORE", "TIPO MEZZO", "TARGA", "COD.MEZZO", "SUPPORTO"
    ]

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h if h else None)
        if h:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    section_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    section_font = Font(name="Calibri", size=10, bold=True, color="000000")

    for _, r in df_spunta.iterrows():
        ridx = int(r["row_idx"])
        is_store = bool(r["is_store"])

        if not is_store:
            # Header line
            ws.cell(row=ridx, column=1, value=r["store_code"]).font = section_font
            ws.cell(row=ridx, column=2, value=r["store_name"]).font = section_font
            ws.cell(row=ridx, column=3, value=r["punto_vendita"]).font = section_font
            for c in range(1, 20):
                ws.cell(row=ridx, column=c).fill = section_fill
            continue

        ws.cell(row=ridx, column=1, value=r["store_code"]).border = thin_border
        ws.cell(row=ridx, column=2, value=r["store_name"]).border = thin_border
        ws.cell(row=ridx, column=3, value=r["punto_vendita"]).border = thin_border
        
        c_mc = ws.cell(row=ridx, column=4, value=round(float(r["mc"]), 2) if r["mc"] > 0 else "")
        c_mc.border = thin_border
        c_mc.number_format = "0.00"

        c_tex = ws.cell(row=ridx, column=5, value=int(r["textile_attesi"]) if r["textile_attesi"] > 0 else "")
        c_tex.border = thin_border

        c_oth = ws.cell(row=ridx, column=7, value=int(r["other_attesi"]) if r["other_attesi"] > 0 else "")
        c_oth.border = thin_border

        c_sho = ws.cell(row=ridx, column=9, value=int(r["shoes_attesi"]) if r["shoes_attesi"] > 0 else "")
        c_sho.border = thin_border

        c_eco = ws.cell(row=ridx, column=11, value=int(r["ecom_attesi"]) if r["ecom_attesi"] > 0 else "")
        c_eco.border = thin_border

        c_ped = ws.cell(row=ridx, column=13, value=int(r["pedane"]) if r["pedane"] > 0 else "")
        c_ped.border = thin_border

        tot_p = r.get("totale_pedane", 0.0)
        c_totp = ws.cell(row=ridx, column=14, value=int(tot_p) if tot_p > 0 else "")
        c_totp.border = thin_border

        ws.cell(row=ridx, column=15, value=r["vettore_autista"]).border = thin_border
        ws.cell(row=ridx, column=16, value=r["tipo_mezzo"]).border = thin_border
        ws.cell(row=ridx, column=17, value=r["targa"]).border = thin_border
        ws.cell(row=ridx, column=18, value=r["cod_mezzo"]).border = thin_border
        ws.cell(row=ridx, column=19, value=r["supporto"]).border = thin_border

        # Giro in col 26 (Z)
        ws.cell(row=ridx, column=26, value=r["giro"])

    # Write Bottom Section: Lavorazioni di Piattaforma (Rows 96-113, dynamically adjusted)
    if special_breakdowns:
        bold_font = Font(name="Calibri", size=10, bold=True)
        title_fill = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
        
        checks = special_breakdowns.get("check_totals", {})
        spec = special_breakdowns.get("special_stores", {})
        zh = special_breakdowns.get("zara_home", {})

        max_spunta_row = int(df_spunta["row_idx"].max()) if not df_spunta.empty and "row_idx" in df_spunta.columns else 75
        base_offset = max(0, (max_spunta_row + 3) - 96)

        # Row 96: Header
        c96 = ws.cell(row=96 + base_offset, column=1, value="TOTALI ")
        c96.font = bold_font
        c96.fill = title_fill
        ws.cell(row=96 + base_offset, column=3, value="CHECK").font = bold_font
        ws.cell(row=96 + base_offset, column=4, value="CHECK").font = bold_font

        # Row 97: Totali Teorici vs Check Spunta Colli
        ws.cell(row=97 + base_offset, column=1, value="TOTALI TEORICI").font = bold_font
        ws.cell(row=97 + base_offset, column=2, value=int(checks.get("totali_teorici_colli", 0))).font = bold_font
        ws.cell(row=97 + base_offset, column=3, value=int(checks.get("totali_check_colli", 0))).font = bold_font

        # Row 98: Totali MC
        ws.cell(row=98 + base_offset, column=1, value="TOTALI MC").font = bold_font
        c_mc_t = ws.cell(row=98 + base_offset, column=2, value=round(float(checks.get("totali_teorici_mc", 0.0)), 2))
        c_mc_t.font = bold_font
        c_mc_t.number_format = "0.00"
        c_mc_c = ws.cell(row=98 + base_offset, column=3, value=round(float(checks.get("totali_check_mc", 0.0)), 2))
        c_mc_c.font = bold_font
        c_mc_c.number_format = "0.00"

        # Row 100: Header Tessile
        ws.cell(row=100 + base_offset, column=2, value="TOTALI").font = bold_font
        ws.cell(row=100 + base_offset, column=3, value="APPESI").font = bold_font
        ws.cell(row=100 + base_offset, column=4, value="STESI").font = bold_font
        ws.cell(row=100 + base_offset, column=10, value="CORREZIONE PROFUMI").font = bold_font

        # Rows 101-103: Tessile Stores
        row_map = {"3680": 101, "9226": 102, "3591": 103}
        for sid, sdata in spec.items():
            r = row_map.get(sid, 101) + base_offset
            ws.cell(row=r, column=1, value=f"PEZZI {sid}").font = bold_font
            ws.cell(row=r, column=2, value=int(sdata.get("totali", 0)))
            ws.cell(row=r, column=3, value=int(sdata.get("appesi", 0)))
            ws.cell(row=r, column=4, value=int(sdata.get("stesi", 0)))
            ws.cell(row=r, column=10, value=int(sdata.get("profumi", 0)))

        # Row 105: Header Zara Home
        ws.cell(row=105 + base_offset, column=1, value="ZARA HOME").font = bold_font
        ws.cell(row=105 + base_offset, column=3, value="/0").font = bold_font
        ws.cell(row=105 + base_offset, column=4, value="/8").font = bold_font

        # Rows 106-107: Zara Home Stores
        zh_row_map = {"1984": 106, "1971": 107}
        for sid, zdata in zh.items():
            r = zh_row_map.get(sid, 106) + base_offset
            ws.cell(row=r, column=1, value=sid).font = bold_font
            ws.cell(row=r, column=2, value=int(zdata.get("totale", 0)))
            ws.cell(row=r, column=3, value=int(zdata.get("barra0", 0)))
            ws.cell(row=r, column=4, value=int(zdata.get("barra9", 0)))

        # Row 108: Totale Pezzi
        ws.cell(row=108 + base_offset, column=1, value="TOTALE PEZZI").font = bold_font
        ws.cell(row=108 + base_offset, column=2, value=int(special_breakdowns.get("totale_pezzi_lavorati", 0))).font = bold_font

        # Row 110: Header Stender / PLT
        ws.cell(row=110 + base_offset, column=1, value="SVILUPPO").font = bold_font
        ws.cell(row=110 + base_offset, column=2, value="STENDER").font = bold_font
        ws.cell(row=110 + base_offset, column=3, value="PLT").font = bold_font

        # Rows 111-113: Sviluppo Stender
        stnd_row_map = {"3680": 111, "9226": 112, "3591": 113}
        for sid, sdata in spec.items():
            r = stnd_row_map.get(sid, 111) + base_offset
            ws.cell(row=r, column=1, value=f"{sid} /0").font = bold_font
            ws.cell(row=r, column=2, value=int(sdata.get("stender", sdata.get("stander", 0))))
            ws.cell(row=r, column=3, value=int(sdata.get("plt", 0)))
            ws.cell(row=r, column=5, value=sdata.get("sviluppo_str", ""))

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col[0].column > 26:
            continue
        max_len = max(len(str(cell.value or "")) for cell in col[:75])
        ws.column_dimensions[col_letter].width = max(max_len + 2, 10)

    wb.save(output_path)


def export_sintesi_operativa(
    df_spunta: pd.DataFrame,
    special_breakdowns: Optional[Dict[str, Any]],
    df_stores: Optional[pd.DataFrame],
    output_path: Union[str, Path],
    run_date: Optional[str] = None
) -> Path:
    """
    Exports an enhanced, professional operational workbook (Sintesi Operativa) with 2 sheets:
    1. Piano Giri & Mezzi (totale colli, volume, pedane, mezzo assegnato, proposta supporto)
    2. Lavorazioni Piattaforma (textile hanging/folded breakdown, perfumes, stander, pallets, Zara Home 50/50)
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # Font definitions
    font_main = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_giro_hdr = Font(name="Calibri", size=11, bold=True, color="0369A1")
    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
    font_tbl_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_tot = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_standby = Font(name="Calibri", size=10, italic=True, color="991B1B")

    # Fill definitions
    fill_title = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_subbar = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    fill_tbl_hdr = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_giro_banner = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    fill_subtot = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_tot = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_standby = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    # Border definitions
    thin_side = Side(style="thin", color="D1D5DB")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Map formulas and info from df_stores
    map_stores = {}
    if df_stores is not None and not df_stores.empty:
        sid_col = "store_id" if "store_id" in df_stores.columns else "store_code"
        for _, s_r in df_stores.iterrows():
            map_stores[str(s_r[sid_col]).strip()] = s_r.to_dict()

    stores_only = df_spunta[df_spunta["is_store"] == True].copy() if not df_spunta.empty else pd.DataFrame()

    tot_colli_all = 0
    tot_vol_all = 0.0
    tot_ped_all = 0

    if not stores_only.empty:
        stores_only["tot_colli_calc"] = (
            stores_only["textile_attesi"].fillna(0) +
            stores_only["other_attesi"].fillna(0) +
            stores_only["shoes_attesi"].fillna(0) +
            stores_only["ecom_attesi"].fillna(0)
        )
        tot_colli_all = int(stores_only["tot_colli_calc"].sum())
        tot_vol_all = round(float(stores_only["mc"].fillna(0).sum()), 2)
        tot_ped_all = int(round(stores_only["pedane"].fillna(0).sum()))

    # -------------------------------------------------------------
    # SHEET 1: PIANO GIRI & MEZZI
    # -------------------------------------------------------------
    ws_giri = wb.active
    ws_giri.title = "Piano Giri & Mezzi"
    ws_giri.views.sheetView[0].showGridLines = True

    # Title Banner (12 columns: A to L)
    ws_giri.merge_cells("A1:L1")
    c_title = ws_giri.cell(row=1, column=1, value="SINTESI OPERATIVA TRASPORTI & GIRI - ZARA LOGISTICS HUB")
    c_title.font = font_title
    c_title.fill = fill_title
    c_title.alignment = align_center
    ws_giri.row_dimensions[1].height = 32

    ws_giri.merge_cells("A2:L2")
    c_sub = ws_giri.cell(
        row=2, column=1,
        value=f"Data Spedizione: {format_date_it(run_date) if run_date else 'N/D'}  |  Totale Colli: {format_number_it(tot_colli_all, 0)}  |  Totale Volume: {format_number_it(tot_vol_all, 2)} m³  |  Totale Pedane: {format_number_it(tot_ped_all, 0)} PLT"
    )
    c_sub.font = font_subtitle
    c_sub.fill = fill_subbar
    c_sub.alignment = align_center
    ws_giri.row_dimensions[2].height = 24

    # Table Header (Row 4) - 12 columns
    tbl_cols = [
        ("Giro", align_center),
        ("Store ID", align_center),
        ("Brand", align_left),
        ("Punto Vendita / Destinazione", align_left),
        ("Totale Colli", align_right),
        ("Volume (m³)", align_right),
        ("Pedane (PLT)", align_right),
        ("Mezzo Assegnato", align_left),
        ("Supporto Mezzo Proposto", align_left),
        ("Vettore", align_left),
        ("Targa", align_center),
        ("Stato Negozio", align_center)
    ]

    curr_row = 4
    for c_i, (col_name, c_align) in enumerate(tbl_cols, 1):
        c = ws_giri.cell(row=curr_row, column=c_i, value=col_name)
        c.fill = fill_tbl_hdr
        c.font = font_tbl_hdr
        c.alignment = c_align
        c.border = thin_border
    ws_giri.row_dimensions[curr_row].height = 24

    def _giro_key(g):
        gs = str(g).strip().upper()
        if gs.startswith("ZA") and gs[2:].isdigit():
            n = int(gs[2:])
            return (1 if 15 <= n <= 19 else 0, n)
        return (2, gs)

    giri_present = sorted(stores_only["giro"].dropna().unique(), key=_giro_key) if not stores_only.empty else []

    curr_row = 5
    for giro in giri_present:
        g_stores = stores_only[stores_only["giro"] == giro].copy()
        if g_stores.empty:
            continue

        first_st = g_stores.iloc[0]
        vett_autista = str(first_st.get("vettore_autista", "") or "")
        tipo_m = str(first_st.get("tipo_mezzo", "") or "")
        targa_m = str(first_st.get("targa", "") or "")
        cod_m = str(first_st.get("cod_mezzo", "") or "")
        supp_m = str(first_st.get("supporto", "") or "")

        base_sid = str(first_st.get("store_code", "")).strip()
        st_meta = map_stores.get(base_sid) or map_stores.get(base_sid[:-1] if base_sid.endswith("B") else base_sid, {})
        vett_str = vett_autista
        if st_meta and st_meta.get("vettore"):
            vett_str = str(st_meta.get("vettore"))

        # Calcolo Proposta Mezzo di Supporto
        g_ped_active = int(sum(
            int(round(float(sr.get("pedane", 0) or 0)))
            for _, sr in g_stores.iterrows()
            if str(sr.get("attivo", map_stores.get(str(sr.get("store_code", "")).strip(), {}).get("attivo", "SI")) or "SI").upper() == "SI"
        ))

        if special_breakdowns and "special_stores" in special_breakdowns:
            for _, sr in g_stores.iterrows():
                sc = str(sr.get("store_code", "")).strip()
                if sc in special_breakdowns["special_stores"] and not bool(sr.get("is_barra9", 0)):
                    sp_info = special_breakdowns["special_stores"][sc]
                    impronta = int(sp_info.get("impronta_a_terra", 0))
                    raw_p = int(round(float(sr.get("pedane", 0) or 0)))
                    if impronta > 0:
                        g_ped_active = g_ped_active - raw_p + impronta

        cap_val = parse_vehicle_capacity(tipo_m)
        supp_res = suggest_support_vehicle(g_ped_active, cap_val)
        if supp_m and supp_m.strip() and supp_m.strip().upper() not in ("NONE", "NAN", "-", ""):
            supp_display = supp_m.strip()
        elif supp_res["status"] == "ok":
            supp_display = "Mezzo Idoneo"
        elif supp_res["status"] == "overload":
            supp_display = f"{supp_res['badge']} ({supp_res['note']})"
        elif supp_res["status"] == "unassigned":
            supp_display = f"{supp_res['badge']} ({supp_res['note']})"
        else:
            supp_display = "—"

        # Giro Banner Row (Cols 1-12)
        ws_giri.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=12)
        banner_text = f"🚛 GIRO: {giro}"
        if vett_str:
            banner_text += f"   |   Vettore: {vett_str}"
        if tipo_m:
            banner_text += f"   |   Mezzo: {tipo_m}"
        if targa_m:
            banner_text += f"   |   Targa: {targa_m}"
        if supp_display and supp_display != "—":
            banner_text += f"   |   Supporto: {supp_display}"

        c_gb = ws_giri.cell(row=curr_row, column=1, value=banner_text)
        c_gb.fill = fill_giro_banner
        c_gb.font = font_giro_hdr
        c_gb.alignment = align_left
        ws_giri.row_dimensions[curr_row].height = 22
        curr_row += 1

        g_colli_tot = 0
        g_vol_tot = 0.0
        g_ped_tot = 0

        for _, s_row in g_stores.iterrows():
            sid = str(s_row.get("store_code", "")).strip()
            bsid = sid[:-1] if sid.endswith("B") else sid
            s_meta = map_stores.get(sid) or map_stores.get(bsid, {})

            tex = int(round(float(s_row.get("textile_attesi", 0) or 0)))
            oth = int(round(float(s_row.get("other_attesi", 0) or 0)))
            sho = int(round(float(s_row.get("shoes_attesi", 0) or 0)))
            eco = int(round(float(s_row.get("ecom_attesi", 0) or 0)))
            c_tot = tex + oth + sho + eco
            vol = round(float(s_row.get("mc", 0.0) or 0.0), 2)
            ped = int(round(float(s_row.get("pedane", 0) or 0)))

            st_att = str(s_row.get("attivo", s_meta.get("attivo", "SI")) or "SI").upper()
            is_op = (st_att == "SI")

            if is_op:
                g_colli_tot += c_tot
                g_vol_tot += vol
                g_ped_tot += ped

            brand_val = str(s_row.get("brand", s_row.get("punto_vendita", "")) or "")
            name_val = str(s_row.get("store_name", "") or "")
            s_vett = str(s_meta.get("vettore", vett_str) or vett_str)
            s_mez = str(s_row.get("tipo_mezzo", tipo_m) or tipo_m)
            s_tar = str(s_row.get("targa", targa_m) or targa_m)
            s_supp = str(s_row.get("supporto", "") or "").strip()
            if not s_supp or s_supp.upper() in ("NONE", "NAN", "-"):
                s_supp = supp_display

            row_data = [
                (giro, align_center, None),
                (sid, align_center, None),
                (brand_val, align_left, None),
                (name_val, align_left, None),
                (c_tot, align_right, "#,##0"),
                (vol, align_right, "0.00"),
                (ped, align_right, "#,##0"),
                (s_mez, align_left, None),
                (s_supp, align_left, None),
                (s_vett, align_left, None),
                (s_tar, align_center, None),
                ("OPERATIVO" if is_op else "STANDBY", align_center, None)
            ]

            for c_i, (v, a_align, n_fmt) in enumerate(row_data, 1):
                cell = ws_giri.cell(row=curr_row, column=c_i, value=v)
                cell.font = font_main if is_op else font_standby
                cell.alignment = a_align
                cell.border = thin_border
                if not is_op:
                    cell.fill = fill_standby
                if n_fmt:
                    cell.number_format = n_fmt

            ws_giri.row_dimensions[curr_row].height = 19
            curr_row += 1

        # Giro Subtotal Row (Cols 1-12)
        subtot_data = [
            (f"SUBTOTALE {giro}", align_left, font_bold),
            ("", align_left, font_bold),
            ("", align_left, font_bold),
            ("", align_left, font_bold),
            (g_colli_tot, align_right, font_bold),
            (round(g_vol_tot, 2), align_right, font_bold),
            (g_ped_tot, align_right, font_bold),
            (tipo_m, align_left, font_bold),
            (supp_display, align_left, font_bold),
            (vett_str, align_left, font_bold),
            (targa_m, align_center, font_bold),
            ("", align_center, font_bold)
        ]

        for c_i, (v, a_align, f_obj) in enumerate(subtot_data, 1):
            cell = ws_giri.cell(row=curr_row, column=c_i, value=v)
            cell.font = f_obj
            cell.fill = fill_subtot
            cell.alignment = a_align
            cell.border = thin_border
            if c_i in [5, 7]:
                cell.number_format = "#,##0"
            elif c_i == 6:
                cell.number_format = "0.00"

        ws_giri.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
        ws_giri.row_dimensions[curr_row].height = 21
        curr_row += 2  # leaves an empty spacer row

    # Grand Total Row at bottom (Cols 1-12)
    tot_data = [
        ("TOTALE GENERALE SPEDIZIONI", align_left),
        ("", align_left),
        ("", align_left),
        ("", align_left),
        (tot_colli_all, align_right),
        (tot_vol_all, align_right),
        (tot_ped_all, align_right),
        ("", align_left),
        ("", align_left),
        ("", align_left),
        ("", align_center),
        ("", align_center)
    ]
    for c_i, (v, a_align) in enumerate(tot_data, 1):
        cell = ws_giri.cell(row=curr_row, column=c_i, value=v)
        cell.font = font_tot
        cell.fill = fill_tot
        cell.alignment = a_align
        cell.border = thin_border
        if c_i in [5, 7]:
            cell.number_format = "#,##0"
        elif c_i == 6:
            cell.number_format = "0.00"

    ws_giri.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
    ws_giri.row_dimensions[curr_row].height = 26

    # Column widths for Sheet 1 (12 columns)
    col_widths_s1 = {
        1: 10, 2: 12, 3: 15, 4: 30, 5: 14, 6: 14, 7: 14,
        8: 18, 9: 25, 10: 18, 11: 14, 12: 14
    }
    for c_idx, w in col_widths_s1.items():
        ws_giri.column_dimensions[get_column_letter(c_idx)].width = w

    # -------------------------------------------------------------
    # SHEET 2: LAVORAZIONI PIATTAFORMA
    # -------------------------------------------------------------
    ws_lav = wb.create_sheet(title="Lavorazioni Piattaforma")
    ws_lav.views.sheetView[0].showGridLines = True

    fill_lav_title = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    fill_lav_subbar = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_lav_hdr = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    fill_lav_subtot = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    font_lav_subbar = Font(name="Calibri", size=10, bold=True, color="065F46")
    font_lav_section = Font(name="Calibri", size=11, bold=True, color="065F46")

    ws_lav.merge_cells("A1:I1")
    c_lt = ws_lav.cell(row=1, column=1, value="PIANO LAVORAZIONI DI PIATTAFORMA CROSS-DOCKING")
    c_lt.font = font_title
    c_lt.fill = fill_lav_title
    c_lt.alignment = align_center
    ws_lav.row_dimensions[1].height = 32

    ws_lav.merge_cells("A2:I2")
    c_ls = ws_lav.cell(row=2, column=1, value="Dettaglio Capi Appesi/Stesi Tessile & Sdoppiamento Zara Home 50/50")
    c_ls.font = font_lav_subbar
    c_ls.fill = fill_lav_subbar
    c_ls.alignment = align_center
    ws_lav.row_dimensions[2].height = 22

    # Section 1: Tessile
    ws_lav.cell(row=4, column=1, value="1. LAVORAZIONE SPECIALE TESSILE (CAPI APPESI, STESI & PROFUMI)").font = font_lav_section

    tex_cols = [
        "Store ID", "Punto Vendita / Negozio", "Totale Capi", "Capi Appesi",
        "Capi Stesi", "Deduzione Profumi", "Stender da Allestire", "Bancali / PLT", "Allestimento Operativo"
    ]
    for c_i, h in enumerate(tex_cols, 1):
        c = ws_lav.cell(row=5, column=c_i, value=h)
        c.fill = fill_lav_hdr
        c.font = font_tbl_hdr
        c.alignment = align_center
        c.border = thin_border
    ws_lav.row_dimensions[5].height = 24

    spec = (special_breakdowns or {}).get("special_stores", {})
    l_row = 6
    sum_capi = 0
    sum_appesi = 0
    sum_stesi = 0
    sum_prof = 0
    sum_stander = 0
    sum_plt = 0

    for sid, sdata in spec.items():
        c_tot = int(sdata.get("totali", sdata.get("totale_pezzi", 0)) or 0)
        c_app = int(sdata.get("appesi", 0) or 0)
        c_ste = int(sdata.get("stesi", 0) or 0)
        c_prf = int(sdata.get("profumi", sdata.get("correzione_profumi", 0)) or 0)
        stnd = int(sdata.get("stender", sdata.get("stander", 0)) or 0)
        banc = int(sdata.get("plt", sdata.get("bancali", 0)) or 0)
        svil = str(sdata.get("sviluppo_str", sdata.get("sintesi", "")) or f"{banc} bancali + {stnd} stender")
        sname = str(sdata.get("store_name", sid))

        sum_capi += c_tot
        sum_appesi += c_app
        sum_stesi += c_ste
        sum_prof += c_prf
        sum_stander += stnd
        sum_plt += banc

        vals = [
            (sid, align_center),
            (sname, align_left),
            (c_tot, align_right),
            (c_app, align_right),
            (c_ste, align_right),
            (c_prf, align_right),
            (stnd, align_right),
            (banc, align_right),
            (svil, align_left)
        ]
        for c_i, (v, a_align) in enumerate(vals, 1):
            cell = ws_lav.cell(row=l_row, column=c_i, value=v)
            cell.font = font_main
            cell.alignment = a_align
            cell.border = thin_border
            if c_i in [3, 4, 5, 6, 7, 8]:
                cell.number_format = "#,##0"

        ws_lav.row_dimensions[l_row].height = 20
        l_row += 1

    # Subtot Tessile
    subtot_tex = [
        ("SUBTOTALE TESSILE", align_left),
        ("", align_left),
        (sum_capi, align_right),
        (sum_appesi, align_right),
        (sum_stesi, align_right),
        (sum_prof, align_right),
        (sum_stander, align_right),
        (sum_plt, align_right),
        (f"{sum_plt} bancali + {sum_stander} stender totali", align_left)
    ]
    for c_i, (v, a_align) in enumerate(subtot_tex, 1):
        cell = ws_lav.cell(row=l_row, column=c_i, value=v)
        cell.font = font_bold
        cell.fill = fill_lav_subtot
        cell.alignment = a_align
        cell.border = thin_border
        if c_i in [3, 4, 5, 6, 7, 8]:
            cell.number_format = "#,##0"

    ws_lav.merge_cells(start_row=l_row, start_column=1, end_row=l_row, end_column=2)
    ws_lav.row_dimensions[l_row].height = 22
    l_row += 3

    # Section 2: Zara Home
    ws_lav.cell(row=l_row-1, column=1, value="2. LAVORAZIONE ZARA HOME (50% BARRA 0 / 50% BARRA 9)").font = font_lav_section
    zh_cols = [
        "Store ID", "Punto Vendita / Negozio", "Totale Pezzi", "Quota Barra 0 (Partenza Diretta)",
        "Quota Barra 9 (Sdoppiamento)", "Note Operative"
    ]
    for c_i, h in enumerate(zh_cols, 1):
        c = ws_lav.cell(row=l_row, column=c_i, value=h)
        c.fill = fill_lav_hdr
        c.font = font_tbl_hdr
        c.alignment = align_center
        c.border = thin_border
    ws_lav.row_dimensions[l_row].height = 24
    l_row += 1

    zh = (special_breakdowns or {}).get("zara_home", {})
    sum_zh_tot = 0
    sum_zh_b0 = 0
    sum_zh_b9 = 0

    for sid, zdata in zh.items():
        ztot = int(zdata.get("totale", zdata.get("totale_paqueteria", 0)) or 0)
        zb0 = int(zdata.get("barra0", zdata.get("quota_barra0", 0)) or 0)
        zb9 = int(zdata.get("barra9", zdata.get("quota_barra8", 0)) or 0)
        zsname = str(zdata.get("store_name", sid))

        sum_zh_tot += ztot
        sum_zh_b0 += zb0
        sum_zh_b9 += zb9

        vals = [
            (sid, align_center),
            (zsname, align_left),
            (ztot, align_right),
            (zb0, align_right),
            (zb9, align_right),
            ("Ripartizione 50% Barra 0 e 50% Barra 9", align_left)
        ]
        for c_i, (v, a_align) in enumerate(vals, 1):
            cell = ws_lav.cell(row=l_row, column=c_i, value=v)
            cell.font = font_main
            cell.alignment = a_align
            cell.border = thin_border
            if c_i in [3, 4, 5]:
                cell.number_format = "#,##0"

        ws_lav.row_dimensions[l_row].height = 20
        l_row += 1

    # Subtot Zara Home
    subtot_zh = [
        ("SUBTOTALE ZARA HOME", align_left),
        ("", align_left),
        (sum_zh_tot, align_right),
        (sum_zh_b0, align_right),
        (sum_zh_b9, align_right),
        ("", align_left)
    ]
    for c_i, (v, a_align) in enumerate(subtot_zh, 1):
        cell = ws_lav.cell(row=l_row, column=c_i, value=v)
        cell.font = font_bold
        cell.fill = fill_lav_subtot
        cell.alignment = a_align
        cell.border = thin_border
        if c_i in [3, 4, 5]:
            cell.number_format = "#,##0"

    ws_lav.merge_cells(start_row=l_row, start_column=1, end_row=l_row, end_column=2)
    ws_lav.row_dimensions[l_row].height = 22
    l_row += 3

    # Section 3: Riepilogo Globale
    ws_lav.cell(row=l_row-1, column=1, value="3. RIEPILOGO TOTALE PIATTAFORMA").font = font_lav_section
    tot_pezzi_lavorati = (special_breakdowns or {}).get("totale_pezzi_lavorati", sum_capi + sum_zh_tot)

    summary_kpis = [
        ("Totale Pezzi / Capi Lavorati in Piattaforma", int(tot_pezzi_lavorati)),
        ("Totale Stender da Allestire (Tessile)", sum_stander),
        ("Totale Bancali da Allestire (Tessile)", sum_plt),
    ]
    for lbl, val in summary_kpis:
        c1 = ws_lav.cell(row=l_row, column=1, value=lbl)
        c1.font = font_bold
        c1.border = thin_border
        c1.fill = fill_lav_subtot
        for col_k in range(2, 4):
            ck = ws_lav.cell(row=l_row, column=col_k, value="")
            ck.border = thin_border
            ck.fill = fill_lav_subtot
        ws_lav.merge_cells(start_row=l_row, start_column=1, end_row=l_row, end_column=3)

        c2 = ws_lav.cell(row=l_row, column=4, value=val)
        c2.font = font_bold
        c2.alignment = align_right
        c2.number_format = "#,##0"
        c2.border = thin_border
        ws_lav.row_dimensions[l_row].height = 21
        l_row += 1

    col_widths_s2 = {
        1: 12, 2: 30, 3: 14, 4: 16, 5: 16, 6: 18, 7: 20, 8: 16, 9: 30
    }
    for c_idx, w in col_widths_s2.items():
        ws_lav.column_dimensions[get_column_letter(c_idx)].width = w

    wb.save(output_path)
    return output_path


def is_cell_gray(cell) -> bool:
    """
    Returns True if the cell has a gray or colored background indicating a non-fillable field.
    Checks openpyxl fill types, RGB values, and theme tints.
    """
    try:
        fill = getattr(cell, "fill", None)
        if not fill or not fill.fill_type or fill.fill_type in ("none", None):
            return False
        sc = getattr(fill, "start_color", None)
        if not sc:
            return False
        if getattr(sc, "type", None) == "rgb" and sc.rgb:
            rgb_str = str(sc.rgb).upper()
            if rgb_str in ("00000000", "FFFFFFFF"):
                return False
            return True
        if getattr(sc, "type", None) == "theme":
            if sc.theme == 0 and (sc.tint is None or sc.tint == 0.0):
                return False
            return True
        return False
    except Exception:
        return False


def export_planner_roma(
    df_planner: pd.DataFrame,
    output_path: Union[str, Path] = None,
    template_path: Optional[Union[str, Path]] = None,
    delivery_date: Optional[Union[str, Any]] = None,
    df_stores: Optional[pd.DataFrame] = None,
    *args,
    **kwargs
) -> Path:
    """
    Exports Planning ITX_RM.xlsx strictly respecting the official Inditex format.
    - Preserves existing sheets ('Planning Roma' and 'DATA ENTRY') and all formulas.
    - Automatically ignores all gray cells ("LE VOCI IN GRIGIO NON VANNO COMPILATE").
    - Populates delivery date in cell D3 (row 3, col 4).
    - Dynamically matches stores and stream types (Textile, Shoes, Barra 9 'B').
    - Handles swapped (template_path, output_path) gracefully.
    """
    # Fallback for extra positional arguments if called like (df, template, output, date)
    if args:
        if delivery_date is None and len(args) >= 1:
            delivery_date = args[0]

    # Handle swapped or default paths
    if output_path is None and template_path is not None:
        output_path = template_path
        template_path = None

    if output_path is not None and template_path is not None:
        p_out = Path(output_path)
        p_tpl = Path(template_path)
        # If output_path is the template and template_path is the output destination
        if p_out.name == "Planning ITX_RM.xlsx" and p_tpl.name != "Planning ITX_RM.xlsx":
            output_path, template_path = template_path, output_path
        elif p_out.exists() and not p_tpl.exists() and "Planning_ITX_RM_" in p_tpl.name:
            output_path, template_path = template_path, output_path

    if template_path is None:
        template_path = PLANNER_TEMPLATE_FILE

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if template_path and Path(template_path).exists():
        wb = openpyxl.load_workbook(template_path)
    else:
        wb = openpyxl.Workbook()

    # Sheet selection - DO NOT rename to PLANNING, as formulas in DATA ENTRY depend on 'Planning Roma'
    if "Planning Roma" in wb.sheetnames:
        ws = wb["Planning Roma"]
    else:
        ws = wb.active

    # 1. Update Date in cell D3 (row 3, col 4) if provided
    if delivery_date:
        try:
            from src.utils import parse_date_obj
            d_obj = parse_date_obj(delivery_date)
            ws.cell(row=3, column=4, value=d_obj.strftime("%d/%m/%Y"))
        except Exception:
            ws.cell(row=3, column=4, value=str(delivery_date).strip())

    # 2. Load store vehicle & driver info from df_stores or database
    if df_stores is None or df_stores.empty:
        try:
            from src.db import get_stores_df
            df_stores = get_stores_df()
        except Exception:
            df_stores = None

    stores_info = {}
    if df_stores is not None and not df_stores.empty:
        sid_col = "store_id" if "store_id" in df_stores.columns else ("store_code" if "store_code" in df_stores.columns else None)
        if sid_col:
            for _, s_row in df_stores.iterrows():
                sid_s = str(s_row[sid_col]).strip()
                stores_info[sid_s] = s_row.to_dict()

    # 3. Index df_planner
    map_by_row_idx = {}
    map_by_code_brand = {}
    map_by_code = {}

    if df_planner is not None and not df_planner.empty:
        for _, r_pl in df_planner.iterrows():
            r_dict = r_pl.to_dict()
            if "row_idx" in r_dict and pd.notna(r_dict["row_idx"]):
                try:
                    map_by_row_idx[int(r_dict["row_idx"])] = r_dict
                except Exception:
                    pass
            sid_s = str(r_dict.get("store_code", r_dict.get("store_id", ""))).strip()
            brand_s = str(r_dict.get("brand", "")).upper()
            is_shoes = "SHOES" in brand_s
            map_by_code_brand[(sid_s, is_shoes)] = r_dict
            if sid_s not in map_by_code:
                map_by_code[sid_s] = r_dict

    def _set_cell(cell, value, is_numeric=True):
        # Rule: LE VOCI IN GRIGIO NON VANNO COMPILATE
        if is_cell_gray(cell):
            return
        # Preserve existing formulas
        if cell.value and str(cell.value).strip().startswith("="):
            return
        if is_numeric:
            try:
                num = float(value)
                if num > 0:
                    cell.value = int(round(num)) if float(num).is_integer() else round(num, 2)
                else:
                    cell.value = None
            except (ValueError, TypeError):
                cell.value = None
        else:
            cell.value = str(value).strip() if value and str(value).strip() else None

    # 4. Iterate dynamically over template data rows (starting at row 7)
    for r_idx in range(7, ws.max_row + 1):
        c_code_val = ws.cell(row=r_idx, column=2).value
        if c_code_val is None:
            continue
        c_code_str = str(c_code_val).strip()
        if not c_code_str or c_code_str.startswith("="):
            continue

        c_brand_val = ws.cell(row=r_idx, column=3).value
        brand_str = str(c_brand_val).strip().upper() if c_brand_val else ""
        is_shoes = "SHOES" in brand_str

        # Match row
        p_row = map_by_row_idx.get(r_idx)
        if not p_row:
            p_row = map_by_code_brand.get((c_code_str, is_shoes))
        if not p_row:
            p_row = map_by_code.get(c_code_str)
        if not p_row:
            continue

        # Col 6 (F): PLANNED CRTS TEXTILE
        _set_cell(ws.cell(row=r_idx, column=6), p_row.get("planned_textile", 0.0))
        # Col 7 (G): RECEIVED CRTS TEXTILE
        _set_cell(ws.cell(row=r_idx, column=7), p_row.get("received_textile", 0.0))
        # Col 8 (H): PLANNED CRTS OTHERS
        _set_cell(ws.cell(row=r_idx, column=8), p_row.get("planned_others", 0.0))
        # Col 9 (I): RECEIVED CRTS OTHERS
        _set_cell(ws.cell(row=r_idx, column=9), p_row.get("received_others", 0.0))
        # Col 10 (J): VOLUME TEXTILE
        _set_cell(ws.cell(row=r_idx, column=10), p_row.get("vol_textile", 0.0))
        # Col 11 (K): WEIGHT TEXTILE
        _set_cell(ws.cell(row=r_idx, column=11), p_row.get("weight_textile", 0.0))

        # Col 12 (L): PLANNED CRTS SHOES
        _set_cell(ws.cell(row=r_idx, column=12), p_row.get("planned_shoes", 0.0))
        # Col 13 (M): RECEIVED CRTS SHOES
        _set_cell(ws.cell(row=r_idx, column=13), p_row.get("received_shoes", 0.0))
        # Col 14 (N): VOLUME SHOES
        _set_cell(ws.cell(row=r_idx, column=14), p_row.get("vol_shoes", 0.0))
        # Col 15 (O): WEIGHT SHOES
        _set_cell(ws.cell(row=r_idx, column=15), p_row.get("weight_shoes", 0.0))

        # Col 16 (P): PLANNED CRTS ECOM
        _set_cell(ws.cell(row=r_idx, column=16), p_row.get("planned_ecom", 0.0))
        # Col 17 (Q): RECEIVED CRTS ECOM
        _set_cell(ws.cell(row=r_idx, column=17), p_row.get("received_ecom", 0.0))
        # Col 18 (R): VOLUME ECOM
        _set_cell(ws.cell(row=r_idx, column=18), p_row.get("vol_ecom", 0.0))
        # Col 19 (S): WEIGHT ECOM
        _set_cell(ws.cell(row=r_idx, column=19), p_row.get("weight_ecom", 0.0))

        # Note: Col 20 (TOTAL VOLUME) & Col 21 (TOTAL WEIGHT) are formulas in template - preserved!

        # Col 22 (V): INTERSHOPS / ID STORE SENDER
        _set_cell(ws.cell(row=r_idx, column=22), p_row.get("intershops", ""), is_numeric=False)

        # Col 29..32: Mezzo, Vettore, Targa, Autista
        st_info = stores_info.get(c_code_str, {})
        # If B store, check base store profile fallback if needed
        if not st_info and c_code_str.endswith("B"):
            st_info = stores_info.get(c_code_str[:-1], {})

        mezzo = p_row.get("tipo_mezzo") or st_info.get("tipo_mezzo")
        vettore = p_row.get("vettore") or st_info.get("vettore")
        targa = p_row.get("targa") or st_info.get("targa")
        autista = p_row.get("autista") or st_info.get("autista")

        if mezzo:
            _set_cell(ws.cell(row=r_idx, column=29), mezzo, is_numeric=False)
        if vettore and not ws.cell(row=r_idx, column=30).value:
            _set_cell(ws.cell(row=r_idx, column=30), vettore, is_numeric=False)
        if targa:
            _set_cell(ws.cell(row=r_idx, column=31), targa, is_numeric=False)
        if autista:
            _set_cell(ws.cell(row=r_idx, column=32), autista, is_numeric=False)

    wb.save(output_path)
    return output_path


def export_bolle_excel(df_bolle: pd.DataFrame, delivery_date: str) -> bytes:
    """Exports delivery notes (bolle) to a styled Excel sheet."""
    from src.bolle_manager import export_bolle_excel as _ebe
    return _ebe(df_bolle, delivery_date)

