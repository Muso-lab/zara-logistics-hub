import pytest
import sqlite3
from datetime import datetime, date, timedelta
import pandas as pd

from src.config import DB_PATH
from src.bolle_manager import (
    get_bolle_dates,
    get_all_drivers,
    get_all_carriers,
    get_store_delivery_profile,
    list_store_delivery_profiles,
    save_store_delivery_profile,
    get_bolle,
    get_bolla_by_id,
    save_or_update_bolla,
    delete_bolla,
    delete_bolle_by_date,
    generate_bolle_from_forecast,
    get_actuals_for_planner,
    render_bolla_printable_html,
    format_intershops_string
)


def test_bolle_crud():
    # Clean up any leftover test data
    existing = get_bolle(data="2099-01-01", id_store="9999")
    for _, row in existing.iterrows():
        delete_bolla(row["id"])

    # 1. Insert a test bolla
    test_date = "2099-01-01"
    b_data = {
        "data": test_date,
        "data_consegna": test_date,
        "id_store": "9999",
        "punto_vendita": "TEST STORE",
        "store_name": "TEST-NEGOZIO",
        "tipo_cons": "STD",
        "vettore": "LACCHI SPA",
        "autista": "MARIO ROSSI",
        "tipo_mezzo": "MOTRICE",
        "indirizzo_destino": "VIA DEI TEST 123",
        "colli_abbigl": 100,
        "colli_calzat": 20,
        "colli_ecom": 10,
        "colli_others": 5,
        "colli_vario": 2,
        "intershops_list": [
            {"mittente": "3680", "colli": 4, "tipo": "CESTE"}
        ]
    }
    b_id = save_or_update_bolla(b_data)
    assert b_id > 0

    # 2. Retrieve
    b_ret = get_bolla_by_id(b_id)
    assert b_ret is not None
    assert b_ret["id_store"] == "9999"
    assert b_ret["colli_abbigl"] == 100.0
    assert "4C CESTE da 3680" in b_ret["stringa"]

    # 3. Update
    b_data["id"] = b_id
    b_data["colli_abbigl"] = 150
    b_id_up = save_or_update_bolla(b_data)
    assert b_id_up == b_id
    b_ret_up = get_bolla_by_id(b_id)
    assert b_ret_up["colli_abbigl"] == 150.0

    # 4. Filter query
    df = get_bolle(data=test_date, id_store="9999")
    assert len(df) == 1
    assert df.iloc[0]["autista"] == "MARIO ROSSI"
    assert "TOT_COLLI" in df.columns
    assert df.iloc[0]["TOT_COLLI"] == 187.0
    assert b_ret_up["TOT_COLLI"] == 187.0

    # 5. Actuals for planner
    df_act = get_actuals_for_planner(test_date)
    assert not df_act.empty
    row_act = df_act[df_act["ID_STORE"] == "9999"].iloc[0]
    assert row_act["TOT_ABBIGL"] == 150.0
    assert row_act["TOT_CALZAT"] == 20.0

    # 6. Render HTML print
    html = render_bolla_printable_html([b_ret_up], title="Test Bolla")
    assert "BOLLA DI CONSEGNA N°" in html
    assert "VIA DEI TEST 123" in html
    assert "MARIO ROSSI" in html

    # 7. Delete
    ok = delete_bolla(b_id)
    assert ok is True
    assert get_bolla_by_id(b_id) is None


def test_intershops_formatting():
    items = [
        {"mittente": "3680", "colli": 6, "tipo": "TRSF"},
        {"mittente": "9226", "colli": 45, "tipo": "CESTE LAVORAZIONE"}
    ]
    s = format_intershops_string(items)
    assert "6C TRSF da 3680" in s
    assert "45C CESTE LAVORAZIONE da 9226" in s


def test_store_delivery_profiles():
    from src.db import delete_store_record
    delete_store_record("8888")
    try:
        prof = {
            "id_store": "8888",
            "tipo_cons": "STD",
            "punto_vendita": "ZARA",
            "store_name": "NEGOZIO PROVA",
            "provincia": "RM",
            "tipo_mezzo": "FURGONE",
            "autista": "GIUSEPPE VERDI",
            "indirizzo": "CORSO ITALIA 1"
        }
        save_store_delivery_profile(prof)
        fetched = get_store_delivery_profile("8888", "STD")
        assert fetched["autista"] == "GIUSEPPE VERDI"
        assert fetched["indirizzo"] == "CORSO ITALIA 1"
    finally:
        delete_store_record("8888")


def test_giro_and_store_registry_management():
    from src.db import (
        get_giro_summary,
        update_giro_driver,
        add_store_record,
        update_store_record,
        get_stores_df,
        delete_store_record,
        get_all_runs
    )

    # 1. Test add_store_record with full driver/vehicle/carrier details
    test_sid = "7777"
    new_store = {
        "store_id": test_sid,
        "store_name": "TEST STORE 7777",
        "punto_vendita": "ZARA",
        "brand": "ZARA",
        "giro": "ZA99",
        "autista": "AUTISTA TEST",
        "cell_autista": "3331122334",
        "tipo_mezzo": "BILICO",
        "vettore": "LACCHI TEST",
        "indirizzo": "VIA DEL TEST 77",
        "provincia": "RM",
        "attivo": "SI",
        "lavorazione": "SI TESSILE",
        "sdoppiamento_b9": "SI",
        "formula": "MC/1,9"
    }
    add_store_record(new_store)

    df_stores = get_stores_df()
    assert test_sid in df_stores["store_id"].values
    store_row = df_stores[df_stores["store_id"] == test_sid].iloc[0]
    assert store_row["giro"] == "ZA99"
    assert store_row["autista"] == "AUTISTA TEST"
    assert store_row["cell_autista"] == "3331122334"
    assert store_row["tipo_mezzo"] == "BILICO"

    # 2. Test get_giro_summary contains ZA99
    df_giri = get_giro_summary()
    assert "ZA99" in df_giri["giro"].values
    g_row = df_giri[df_giri["giro"] == "ZA99"].iloc[0]
    assert g_row["autista"] == "AUTISTA TEST"
    assert g_row["tipo_mezzo"] == "BILICO"

    # 3. Test update_giro_driver bulk update
    n_up = update_giro_driver("ZA99", autista="NUOVO AUTISTA", cell_autista="3339999999", tipo_mezzo="MOTRICE 18 PLT", vettore="LACCHI SPA")
    assert n_up >= 1
    df_stores2 = get_stores_df()
    store_row2 = df_stores2[df_stores2["store_id"] == test_sid].iloc[0]
    assert store_row2["autista"] == "NUOVO AUTISTA"
    assert store_row2["cell_autista"] == "3339999999"
    assert store_row2["tipo_mezzo"] == "MOTRICE 18 PLT"

    # 4. Test get_store_delivery_profile returns updated driver/cell/mezzo
    prof = get_store_delivery_profile(test_sid, "STD")
    assert prof["autista"] == "NUOVO AUTISTA"
    assert prof["cell_autista"] == "3339999999"
    assert prof["tipo_mezzo"] == "MOTRICE 18 PLT"

    # 5. Test get_all_runs has total_planned_crts
    runs = get_all_runs()
    if not runs.empty:
        assert "total_planned_crts" in runs.columns
        assert "total_bultos" in runs.columns

    # Clean up
    delete_store_record(test_sid)


def test_generate_bolle_and_consuntivo_planner():
    from src.bolle_manager import generate_bolle_from_forecast, get_actuals_for_planner
    from src.db import get_all_runs, get_run_details
    from src.engine import compute_planning_roma

    runs = get_all_runs()
    if runs.empty:
        return

    test_run_id = int(runs.iloc[0]["id"])
    test_delivery_date = "2099-12-31"

    # 1. Test generate_bolle_from_forecast
    created_n, skipped_n = generate_bolle_from_forecast(
        run_id=test_run_id,
        delivery_date=test_delivery_date,
        overwrite=True
    )
    assert created_n > 0

    # 2. Test get_actuals_for_planner
    df_actuals = get_actuals_for_planner(test_delivery_date)
    assert not df_actuals.empty
    assert "ID_STORE" in df_actuals.columns
    assert "TOT_ABBIGL" in df_actuals.columns

    # 3. Test compute_planning_roma with df_a_target (which has store_code, NOT store_id!)
    details = get_run_details(test_run_id)
    df_planning = details.get("planning")
    assert df_planning is not None and not df_planning.empty
    assert "store_code" in df_planning.columns
    # Crucial test: must not raise KeyError: "None of ['store_id'] are in the columns"
    df_planner_res = compute_planning_roma(df_planning, df_actuals)
    assert not df_planner_res.empty
    assert "total_planned_crts" in df_planner_res.columns
    assert "total_received_crts" in df_planner_res.columns

    # Clean up test bolle
    delete_bolle_by_date(test_delivery_date)
    assert df_planner_res["total_planned_crts"].sum() > 0


def test_export_bolle_excel():
    import io
    import openpyxl
    from src.bolle_manager import export_bolle_excel

    sample_bolle = pd.DataFrame([
        {
            "giro": "ZA01",
            "id_store": "3680",
            "punto_vendita": "ZARA",
            "store_name": "ROMA PORTA DI ROMA",
            "tipo_cons": "STD",
            "colli_abbigl": 150,
            "colli_calzat": 30,
            "colli_ecom": 15,
            "colli_others": 5,
            "colli_vario": 2,
            "vettore": "LACCHI SPA",
            "tipo_mezzo": "MOTRICE 18 PLT",
            "indirizzo_destino": "VIA ALBERTO LIONELLO 201",
            "stringa": "4C CESTE",
            "anomalie_scarico": ""
        },
        {
            "giro": "ZA01",
            "id_store": "3680",
            "punto_vendita": "ZARA",
            "store_name": "ROMA PORTA DI ROMA",
            "tipo_cons": "DOPPIA",
            "colli_abbigl": 65,
            "colli_calzat": 12,
            "colli_ecom": 5,
            "colli_others": 2,
            "colli_vario": 0,
            "vettore": "LACCHI SPA",
            "tipo_mezzo": "FURGONE",
            "indirizzo_destino": "VIA ALBERTO LIONELLO 201",
            "stringa": "",
            "anomalie_scarico": ""
        }
    ])

    excel_bytes = export_bolle_excel(sample_bolle, "2026-10-27")
    assert excel_bytes is not None and len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Riepilogo Bolle" in wb.sheetnames
    ws = wb["Riepilogo Bolle"]
    # Check Title
    assert "RIEPILOGO BOLLE DI CONSEGNA INDITEX" in ws["A1"].value
    # Check Headers
    assert ws["A2"].value == "Giro"
    assert ws["F2"].value == "Totale Colli"
    # Check Rows (row 3, 4 data, row 5 total)
    assert ws["B3"].value == "3680"
    assert ws["E3"].value == "STD"
    assert ws["F3"].value == 202
    assert ws["E4"].value == "DOPPIA"
    assert ws["F4"].value == 84
    assert ws["A5"].value == "TOTALE"
    assert "=SUM(F3:F4)" in ws["F5"].value


def test_bolle_manager_db_path_parameter():
    from src.bolle_manager import get_actuals_for_planner, get_bolle, get_bolle_dates
    # Verify db_path keyword argument is accepted without error
    df_act = get_actuals_for_planner("2099-01-01", db_path=DB_PATH)
    assert isinstance(df_act, pd.DataFrame)
    df_b = get_bolle(data="2099-01-01", db_path=DB_PATH)
    assert isinstance(df_b, pd.DataFrame)
    dates = get_bolle_dates(5, db_path=DB_PATH)
    assert isinstance(dates, list)

