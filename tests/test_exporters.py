import tempfile
from pathlib import Path
import openpyxl
import pandas as pd
import pytest

from src.engine import compute_sga_flows
from src.exporters import export_sga_file, export_sintesi_operativa


def test_sga_flows_and_export():
    # Setup sample spunta dataframe
    df_spunta = pd.DataFrame([
        {
            "row_idx": 1,
            "is_store": True,
            "store_code": "3503",
            "store_name": "ROMA VIA CONDOTTI",
            "punto_vendita": "ZARA",
            "brand": "ZARA",
            "mc": 4.5,
            "textile_attesi": 100,
            "other_attesi": 20,
            "shoes_attesi": 15,
            "ecom_attesi": 5,
            "pedane": 6,
            "totale_pedane": 6,
            "vettore_autista": "MARIO ROSSI",
            "tipo_mezzo": "EUROCARGO 75",
            "targa": "AA123BB",
            "cod_mezzo": "EC75",
            "supporto": "",
            "giro": "ZA1",
            "is_barra9": 0,
            "attivo": "SI"
        },
        {
            "row_idx": 2,
            "is_store": True,
            "store_code": "9999",
            "store_name": "NEGOZIO NUOVO NON ASSEGNATO",
            "punto_vendita": "ZARA",
            "brand": "ZARA",
            "mc": 2.0,
            "textile_attesi": 50,
            "other_attesi": 0,
            "shoes_attesi": 0,
            "ecom_attesi": 0,
            "pedane": 3,
            "totale_pedane": 3,
            "vettore_autista": "",
            "tipo_mezzo": "",
            "targa": "",
            "cod_mezzo": "",
            "supporto": "",
            "giro": "DA ASSEGNARE",
            "is_barra9": 0,
            "attivo": "SI"
        },
        {
            "row_idx": 3,
            "is_store": True,
            "store_code": "1984B",
            "store_name": "ROMA COLA DI RIENZO",
            "punto_vendita": "ZARA HOME",
            "brand": "ZARA HOME",
            "mc": 3.0,
            "textile_attesi": 0,
            "other_attesi": 0,
            "shoes_attesi": 0,
            "ecom_attesi": 0,
            "pedane": 4,
            "totale_pedane": 4,
            "vettore_autista": "LUIGI BIANCHI",
            "tipo_mezzo": "MOTRICE 18 PLT",
            "targa": "CC456DD",
            "cod_mezzo": "MOT18",
            "supporto": "",
            "giro": "ZA15",
            "is_barra9": 1,
            "attivo": "SI"
        },
        {
            "row_idx": 4,
            "is_store": True,
            "store_code": "8888",
            "store_name": "NEGOZIO B9 DA ASSEGNARE",
            "punto_vendita": "ZARA",
            "brand": "ZARA",
            "mc": 2.0,
            "textile_attesi": 0,
            "other_attesi": 0,
            "shoes_attesi": 0,
            "ecom_attesi": 0,
            "pedane": 2,
            "totale_pedane": 2,
            "vettore_autista": "",
            "tipo_mezzo": "",
            "targa": "",
            "cod_mezzo": "",
            "supporto": "",
            "giro": "DA ASSEGNARE",
            "is_barra9": 1,
            "attivo": "SI"
        }
    ])

    df_stores = pd.DataFrame([
        {"store_id": "3503", "cod_sga": 3503, "giro": "ZA1", "attivo": "SI"},
        {"store_id": "9999", "cod_sga": 9999, "giro": "DA ASSEGNARE", "attivo": "SI"},
        {"store_id": "1984", "cod_sga": 1984, "giro": "ZA15", "attivo": "SI"},
        {"store_id": "8888", "cod_sga": 8888, "giro": "DA ASSEGNARE", "attivo": "SI"}
    ])

    df_sga0, df_sga9 = compute_sga_flows(df_spunta, df_stores, "2026-10-26")

    # 1. Unassigned stores (9999 and 8888) MUST NOT be present in SGA flows
    assert "9999" not in df_sga0["store_code"].values
    assert "8888" not in df_sga9["store_code"].values
    assert (df_sga0["GIRO"] == "DA ASSEGNARE").sum() == 0
    assert (df_sga9["GIRO"] == "DA ASSEGNARE").sum() == 0

    # 2. Date format MUST be DD/MM/YYYY (gg/mm/aaaa)
    assert df_sga0.iloc[0]["DATA_CARICO"] == "27/10/2026"
    assert df_sga9.iloc[0]["DATA_CARICO"] == "28/10/2026"

    # 3. Test export to Excel and CSV
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_p0 = Path(tmp_dir) / "barrazero_Sga.xlsx"
        export_sga_file(df_sga0, tmp_p0)

        # Check Excel content
        wb0 = openpyxl.load_workbook(tmp_p0)
        ws0 = wb0["SGA_DEF"]
        # Date cell check (row 2, col 2 is DATA_CARICO)
        date_val = str(ws0.cell(row=2, column=2).value)
        assert date_val == "27/10/2026"

        # Check CSV content
        csv0_p = tmp_p0.with_suffix(".csv")
        assert csv0_p.exists()
        csv_df = pd.read_csv(csv0_p)
        assert str(csv_df.iloc[0]["DATA_CARICO"]) == "27/10/2026"
        assert "DA ASSEGNARE" not in csv_df["GIRO"].values


def test_sintesi_operativa_structure():
    df_spunta = pd.DataFrame([
        {
            "row_idx": 1,
            "is_store": True,
            "store_code": "3503",
            "store_name": "ROMA VIA CONDOTTI",
            "punto_vendita": "ZARA",
            "brand": "ZARA",
            "mc": 12.5,
            "textile_attesi": 200,
            "other_attesi": 50,
            "shoes_attesi": 30,
            "ecom_attesi": 10,
            "pedane": 18,
            "totale_pedane": 18,
            "vettore_autista": "MARIO ROSSI",
            "tipo_mezzo": "EUROCARGO 75",  # Capienza 14 PLT -> Overload with 18 PLT!
            "targa": "AA123BB",
            "cod_mezzo": "EC75",
            "supporto": "",
            "giro": "ZA1",
            "is_barra9": 0,
            "attivo": "SI"
        }
    ])

    special_breakdowns = {
        "special_stores": {
            "3680": {
                "store_name": "C.C. Porta di Roma",
                "totale_pezzi": 3000,
                "appesi": 1200,
                "stesi": 1800,
                "correzione_profumi": 1085,
                "stander": 6,
                "bancali": 6,
                "sintesi": "6 bancali + 6 stander"
            }
        },
        "zara_home": {
            "1984": {
                "store_name": "Roma Cola di Rienzo",
                "totale_paqueteria": 2000,
                "quota_barra0": 1000,
                "quota_barra8": 1000
            }
        },
        "totale_pezzi_lavorati": 5000
    }

    df_stores = pd.DataFrame([
        {"store_id": "3503", "giro": "ZA1", "attivo": "SI", "tipo_mezzo": "EUROCARGO 75", "vettore": "VETTORE X", "autista": "MARIO"}
    ])

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_sintesi = Path(tmp_dir) / "Sintesi_Operativa.xlsx"
        export_sintesi_operativa(df_spunta, special_breakdowns, df_stores, tmp_sintesi, run_date="2026-10-26")

        wb = openpyxl.load_workbook(tmp_sintesi)
        sheet_names = wb.sheetnames

        # Sheet 1 and 2 must exist, Sheet 3 MUST NOT exist
        assert "Piano Giri & Mezzi" in sheet_names
        assert "Lavorazioni Piattaforma" in sheet_names
        assert "Elenco Punti Vendita" not in sheet_names
        assert len(sheet_names) == 2

        # Check Sheet 1 headers (Row 4 - 12 columns, Autista removed)
        ws1 = wb["Piano Giri & Mezzi"]
        headers = [ws1.cell(row=4, column=c).value for c in range(1, 13)]
        assert headers == [
            "Giro",
            "Store ID",
            "Brand",
            "Punto Vendita / Destinazione",
            "Totale Colli",
            "Volume (m³)",
            "Pedane (PLT)",
            "Mezzo Assegnato",
            "Supporto Mezzo Proposto",
            "Vettore",
            "Targa",
            "Stato Negozio"
        ]

        # Check Sheet 2 header for Stender
        ws2 = wb["Lavorazioni Piattaforma"]
        assert ws2.cell(row=5, column=7).value == "Stender da Allestire"

        # Check support vehicle proposal for overloaded Giro ZA1 (18 PLT on Eurocargo 14 PLT)
        # Row 5 is Giro Banner, Row 6 is Store 3503
        supp_cell = ws1.cell(row=6, column=9).value
        assert supp_cell is not None and "FURGONE" in str(supp_cell)
